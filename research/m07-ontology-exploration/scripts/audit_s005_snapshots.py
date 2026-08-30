#!/usr/bin/env python3
"""Read-only structural audit for S005 investment snapshot workbooks."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


SNAPSHOT_RE = re.compile(r"(20\d{6})")


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized_cells(row: tuple[Any, ...]) -> list[Any]:
    return [json_value(value) for value in row]


def header_score(row: tuple[Any, ...]) -> tuple[int, int]:
    values = [value for value in row if value not in (None, "")]
    text_values = [value for value in values if isinstance(value, str)]
    return len(text_values), len(values)


def detect_header(worksheet, scan_rows: int = 25) -> tuple[int, list[Any]]:
    candidates: list[tuple[tuple[int, int], int, tuple[Any, ...]]] = []
    for index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(worksheet.max_row, scan_rows), values_only=True),
        start=1,
    ):
        candidates.append((header_score(row), index, row))
    if not candidates:
        return 1, []
    _, index, row = max(candidates, key=lambda item: (item[0][0], item[0][1], -item[1]))
    return index, normalized_cells(row)


def inspect_workbook(path: Path, sample_rows: int) -> dict[str, Any]:
    snapshot_match = SNAPSHOT_RE.search(path.name)
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheets: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        header_row, headers = detect_header(worksheet)
        samples: list[list[Any]] = []
        if sample_rows:
            start = min(header_row + 1, worksheet.max_row)
            end = min(worksheet.max_row, header_row + sample_rows)
            for row in worksheet.iter_rows(min_row=start, max_row=end, values_only=True):
                samples.append(normalized_cells(row))
        sheets.append(
            {
                "name": worksheet.title,
                "rows": worksheet.max_row,
                "columns": worksheet.max_column,
                "headerRow": header_row,
                "headers": headers,
                "sampleRows": samples,
            }
        )
    workbook.close()
    return {
        "file": path.name,
        "path": str(path),
        "snapshotDate": snapshot_match.group(1) if snapshot_match else None,
        "sha256": file_sha256(path),
        "sizeBytes": path.stat().st_size,
        "sheets": sheets,
    }


def summarize(workbooks: list[dict[str, Any]]) -> dict[str, Any]:
    signature_counts: Counter[str] = Counter()
    sheet_counts: Counter[str] = Counter()
    for workbook in workbooks:
        for sheet in workbook["sheets"]:
            sheet_counts[sheet["name"]] += 1
            signature = json.dumps(sheet["headers"], ensure_ascii=False, sort_keys=True)
            signature_counts[signature] += 1
    dates = [workbook["snapshotDate"] for workbook in workbooks if workbook["snapshotDate"]]
    return {
        "schemaVersion": "ofw.m07.s005.snapshot-audit.v1",
        "workbookCount": len(workbooks),
        "dateRange": {"from": min(dates) if dates else None, "to": max(dates) if dates else None},
        "sheetNameCounts": dict(sheet_counts),
        "headerSignatureCount": len(signature_counts),
        "headerSignatures": [
            {"count": count, "headers": json.loads(signature)}
            for signature, count in signature_counts.most_common()
        ],
        "workbooks": workbooks,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("snapshot_dir", type=Path)
    parser.add_argument("--file", action="append", default=[], help="Inspect only the named workbook")
    parser.add_argument("--sample-rows", type=int, default=0)
    parser.add_argument("--json-out", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    paths = sorted(args.snapshot_dir.glob("*.xlsx"))
    if args.file:
        wanted = set(args.file)
        paths = [path for path in paths if path.name in wanted]
    workbooks = [inspect_workbook(path, max(0, args.sample_rows)) for path in paths]
    result = summarize(workbooks)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_out:
        args.json_out.parent.mkdir(parents=True, exist_ok=True)
        args.json_out.write_text(payload + "\n", encoding="utf-8")
    else:
        print(payload)


if __name__ == "__main__":
    main()
