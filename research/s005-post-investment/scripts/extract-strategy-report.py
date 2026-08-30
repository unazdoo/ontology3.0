#!/usr/bin/env python3
"""Read-only structured inspection for the S005 strategy report DOCX."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def paragraph_record(index: int, paragraph) -> dict:
    text = paragraph.text.strip()
    return {
        "index": index,
        "style": paragraph.style.name if paragraph.style else None,
        "text": text,
        "character_count": len(text),
    }


def table_record(index: int, table, include_cells: bool) -> dict:
    rows = []
    if include_cells:
        rows = [
            [cell.text.strip() for cell in row.cells]
            for row in table.rows
        ]
    return {
        "index": index,
        "row_count": len(table.rows),
        "column_count": max((len(row.cells) for row in table.rows), default=0),
        "rows": rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("path", type=Path)
    parser.add_argument(
        "--mode",
        choices=("summary", "outline", "paragraphs", "tables", "all"),
        default="summary",
    )
    parser.add_argument("--start", type=int, default=0)
    parser.add_argument("--end", type=int)
    args = parser.parse_args()

    path = args.path.resolve(strict=True)
    document = Document(path)
    paragraphs = [
        paragraph_record(index, paragraph)
        for index, paragraph in enumerate(document.paragraphs)
    ]
    nonempty = [record for record in paragraphs if record["text"]]
    headings = [
        record
        for record in nonempty
        if (record["style"] or "").lower().startswith("heading")
        or "title" in (record["style"] or "").lower()
    ]
    tables = [
        table_record(index, table, args.mode in {"tables", "all"})
        for index, table in enumerate(document.tables)
    ]

    result = {
        "source": {
            "path": str(path),
            "size_bytes": path.stat().st_size,
            "sha256": sha256(path),
        },
        "structure": {
            "paragraph_count": len(paragraphs),
            "nonempty_paragraph_count": len(nonempty),
            "heading_count": len(headings),
            "table_count": len(tables),
            "table_shapes": [
                [table["row_count"], table["column_count"]] for table in tables
            ],
        },
    }

    if args.mode == "outline":
        result["headings"] = headings
    elif args.mode in {"paragraphs", "all"}:
        result["paragraphs"] = paragraphs[args.start : args.end]
    if args.mode in {"tables", "all"}:
        result["tables"] = tables

    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
