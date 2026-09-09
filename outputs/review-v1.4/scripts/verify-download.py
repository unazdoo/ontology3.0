import datetime
import hashlib
import json
from decimal import Decimal
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[3]
out = root / 'outputs/review-v1.4'
source = root / 'designs/prototype-work/v1.4/public/data'
catalog = json.loads((source / 'snapshots/catalog.json').read_text())
data = json.loads((source / 'portfolio.json').read_text())
file = out / 'downloads/finance-risk-snapshot.xlsx'
workbook = openpyxl.load_workbook(file, data_only=True)
errors, checks = [], []
count = 0
for asset in catalog['assets']:
    sheet = workbook[asset['sheetName']]
    rows = list(sheet.iter_rows(min_row=2))
    count += len(rows)
    if len(rows) != len(asset['rows']):
        errors.append(f"{sheet.title}: row count mismatch")
    for ri, (cells, expected) in enumerate(zip(rows, asset['rows']), 2):
        for ci, (cell, field) in enumerate(zip(cells, asset['fields']), 1):
            want = expected[field['key']]
            actual = cell.value
            if field['type'] == 'percent':
                good = isinstance(actual, (int, float)) and abs(actual * 100 - want) < 1e-9 and '%' in cell.number_format
            elif field['type'] == 'date':
                good = isinstance(actual, (datetime.date, datetime.datetime)) and actual.strftime('%Y-%m-%d') == want
            elif isinstance(want, (int, float)):
                good = isinstance(actual, (int, float)) and abs(actual - want) < 1e-8
            else:
                good = (actual or '') == (want or '')
            if not good:
                errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'actual': str(actual), 'expected': want})
    checks.append({'asset': asset['id'], 'sheet': sheet.title, 'rows': len(rows), 'fields': len(asset['fields']), 'freeze': str(sheet.freeze_panes)})

pages = json.loads((out / 'logs/loan-table-all-pages.json').read_text())
ui_rows = [r for page in pages for r in page]
loans = {r['id']: r for r in data['loans']}
for row in ui_rows:
    loan = loans[row[0]]
    if row[1] != loan['enterpriseId'] or abs(float(row[6].replace(',', '')) - loan['principal']) > 1e-8 or abs(float(row[7])-loan['rate']) > 1e-8 or row[9] != loan['startDate'] or row[10] != loan['maturityDate']:
        errors.append({'uiRow': row, 'source': loan})
if set(r[0] for r in ui_rows) != set(loans):
    errors.append('UI loan identities do not match source')

D = lambda x: Decimal(str(x))
balance = sum(D(l['principal']) for l in loans.values())
interest = sum(D(l['principal']) * D(l['rate']) / 100 for l in loans.values())
end = '2026-03-31'
due = sum(D(l['principal']) for l in loans.values() if data['asOf'] < l['maturityDate'] <= end)
gap = Decimal(0)
for entity in data['enterprises']:
    entity_due = sum(D(l['principal']) for l in loans.values() if l['enterpriseId'] == entity['id'] and data['asOf'] < l['maturityDate'] <= end)
    credit = sum(D(f['undrawn']) for f in data['facilities'] if f['enterpriseId'] == entity['id'] and f['validUntil'] >= end)
    gap += max(Decimal(0), entity_due - D(entity['cash']) - credit)
sha = hashlib.sha256(file.read_bytes()).hexdigest()
result = {'file': str(file), 'sha256': sha, 'byteLength': file.stat().st_size, 'catalogFile': catalog['file'], 'totalRecords': count, 'worksheets': len(workbook.sheetnames), 'sheets': checks, 'uiLoanPages': list(map(len,pages)), 'uiLoanIDs': len(set(r[0] for r in ui_rows)), 'independentBaseline': {'principalMillions': str(balance), 'costPercent': str(interest / balance * 100), 'due90Millions': str(due), 'gap90Millions': str(gap)}, 'errors': errors, 'retryBytesIdentical': file.read_bytes() == (out/'downloads/finance-risk-retry.xlsx').read_bytes()}
(out/'logs/download-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
